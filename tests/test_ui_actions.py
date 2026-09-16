"""Manual reviews and legacy uploads are explicit, separate user actions."""

import base64
import copy
from io import BytesIO

import duckdb
import pytest
from openpyxl import load_workbook

from roadinsight_ui.actions import (
    MAX_UPLOAD_BYTES,
    read_reviews,
    record_review,
    save_upload,
    validate_upload,
)
from roadinsight_ui.snapshot import load_snapshot


def test_manual_review_is_separate_idempotent_and_version_bound(tmp_path):
    snapshot = load_snapshot()
    original = copy.deepcopy(snapshot)
    issue = snapshot["issues"][0]["issue_id"]
    first = record_review(
        snapshot,
        issue,
        "needs_review",
        "测试核验人",
        "仍需核对现场通行标识",
        tmp_path,
        "event-1",
    )
    record_review(
        snapshot,
        issue,
        "needs_review",
        "测试核验人",
        "仍需核对现场通行标识",
        tmp_path,
        "event-1",
    )
    rows = read_reviews(snapshot, tmp_path)
    assert len(rows) == 1 and rows[0]["review_id"] == first["review_id"]
    assert rows[0]["analysis_run_id"] == snapshot["analysis_run"]["run_id"]
    assert snapshot == original
    snapshot["analysis_run"]["run_id"] = "different-run"
    assert read_reviews(snapshot, tmp_path) == []


@pytest.mark.parametrize(
    "issue,conclusion,reviewer,note",
    [
        ("unknown", "confirmed", "测试人", "已检查现场规则"),
        (None, "auto_verified", "测试人", "已检查现场规则"),
        (None, "confirmed", "", "已检查现场规则"),
        (None, "confirmed", "测试人", "好"),
    ],
)
def test_invalid_review_is_rejected(tmp_path, issue, conclusion, reviewer, note):
    snapshot = load_snapshot()
    with pytest.raises(ValueError):
        record_review(
            snapshot,
            issue or snapshot["issues"][0]["issue_id"],
            conclusion,
            reviewer,
            note,
            tmp_path,
            "event",
        )
    assert read_reviews(snapshot, tmp_path) == []


def test_legacy_upload_validation_storage_and_excel_report(tmp_path):
    content = "record_id,report_time,longitude,latitude,issue_type,source,description,region\nA1,2026-09-16 10:00,121.43,31.18,missing_road,demo,疑似缺路,徐汇\nA2,2026-09-16 10:00,999,31.18,missing_road,demo,坐标异常,徐汇\n".encode()
    upload = validate_upload("observations.csv", base64.b64encode(content).decode())
    assert upload["preview"]["summary"]["valid_records"] == 1
    assert upload["preview"]["summary"]["rejected_records"] == 1
    result = save_upload(upload, tmp_path / "uploads.duckdb")
    with duckdb.connect(str(tmp_path / "uploads.duckdb"), read_only=True) as db:
        assert result["batch_id"]
        assert len(db.execute("SHOW TABLES").fetchall()) >= 4
    workbook = load_workbook(
        BytesIO(base64.b64decode(result["report_base64"])), read_only=True
    )
    assert len(workbook.sheetnames) == 5
    workbook.close()


def test_upload_limits_and_missing_columns_fail_without_persistence(tmp_path):
    with pytest.raises(ValueError):
        validate_upload("bad.csv", "!")
    with pytest.raises(ValueError, match="5 MB"):
        validate_upload("big.csv", "a" * (MAX_UPLOAD_BYTES * 2))
    with pytest.raises(ValueError):
        validate_upload("bad.exe", base64.b64encode(b"hello").decode())
    upload = validate_upload(
        "missing.csv", base64.b64encode(b"record_id\nA1\n").decode()
    )
    assert upload["preview"]["summary"]["missing_required_columns"]
    with pytest.raises(ValueError, match="必填列"):
        save_upload(upload, tmp_path / "uploads.duckdb")
    assert not (tmp_path / "uploads.duckdb").exists()
