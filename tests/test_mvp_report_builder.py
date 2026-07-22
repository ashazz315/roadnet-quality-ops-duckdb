from __future__ import annotations

import unittest
from datetime import datetime
from io import BytesIO

import pandas as pd
from openpyxl import load_workbook

from quality_mvp.report_builder import build_daily_report, build_error_csv, daily_report_name


class ReportBuilderTests(unittest.TestCase):
    def test_error_csv_has_utf8_bom_and_error_fields(self) -> None:
        errors = pd.DataFrame([{"record_id": "A1", "error_messages": "经度必须是数字"}])
        payload = build_error_csv(errors)
        self.assertTrue(payload.startswith(b"\xef\xbb\xbf"))
        self.assertIn("经度必须是数字", payload.decode("utf-8-sig"))

    def test_daily_report_can_be_reopened(self) -> None:
        generated_at = datetime(2026, 7, 22, 9, 30)
        valid = pd.DataFrame([{"record_id": "A1", "longitude": 116.4, "latitude": 39.9}])
        errors = pd.DataFrame([{"record_id": "A2", "error_messages": "纬度超出范围"}])
        payload = build_daily_report(valid, errors, "sample.csv", generated_at)
        workbook = load_workbook(BytesIO(payload), read_only=True, data_only=True)
        self.assertEqual(workbook.sheetnames, ["日报汇总", "有效记录", "错误记录"])
        self.assertEqual(workbook["日报汇总"]["B4"].value, 2)
        self.assertEqual(workbook["日报汇总"]["B5"].value, 1)
        self.assertEqual(workbook["日报汇总"]["B6"].value, 1)
        self.assertEqual(daily_report_name(generated_at).name, "road_quality_daily_20260722.xlsx")


if __name__ == "__main__":
    unittest.main()
