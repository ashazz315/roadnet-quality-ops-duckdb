from __future__ import annotations

from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook

from src.daily_report import REPORT_SHEETS, build_daily_report_excel, load_report_queries, query_daily_report
from src.storage import persist_batch
from tests.test_storage import sample_frames


def test_named_sql_queries_return_batch_data_instead_of_demo_numbers(tmp_path: Path) -> None:
    database = tmp_path / "quality.duckdb"
    valid, rejected, summary = sample_frames()
    batch_id = persist_batch(
        "actual_upload.xlsx",
        valid,
        rejected,
        summary,
        database_path=database,
        uploaded_at=datetime(2026, 7, 22, tzinfo=timezone.utc),
    )

    queries = load_report_queries()
    assert set(REPORT_SHEETS) <= set(queries)
    report = query_daily_report(batch_id, database)
    overview = report["data_overview"].iloc[0]
    assert overview["数据源"] == "actual_upload.xlsx"
    assert int(overview["总记录数"]) == len(valid) + len(rejected)
    assert int(overview["有效记录数"]) == len(valid)
    assert int(overview["错误记录数"]) == len(rejected)
    assert int(overview["问题工单数"]) == 1
    assert report["issue_type_statistics"].iloc[0]["问题类型"] == "point_offset"
    assert set(report["region_statistics"]["区域"]) == {"朝阳区", "海淀区"}
    assert report["error_details"].iloc[0]["longitude"] == "bad"
    assert "必须保留" in report["error_details"].iloc[0]["原始记录(JSON)"]


def test_excel_report_has_all_required_sheets_and_is_reopenable(tmp_path: Path) -> None:
    database = tmp_path / "quality.duckdb"
    valid, rejected, summary = sample_frames()
    batch_id = persist_batch("input.csv", valid, rejected, summary, database_path=database)
    payload = build_daily_report_excel(query_daily_report(batch_id, database))

    workbook = load_workbook(BytesIO(payload), read_only=True, data_only=True)
    assert workbook.sheetnames == list(REPORT_SHEETS.values())
    assert workbook["数据概览"]["C2"].value == 3
    assert workbook["高优先级问题"]["A2"].value == "A1"
    assert workbook["错误数据明细"]["A2"].value == "A3"
    workbook.close()


def test_streamlit_entrypoint_does_not_group_metrics_with_pandas() -> None:
    app_source = (Path(__file__).resolve().parents[1] / "app.py").read_text(encoding="utf-8")
    assert ".groupby(" not in app_source
    assert "query_daily_report" in app_source
    assert "persist_batch" in app_source
