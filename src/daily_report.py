"""SQL-backed daily report queries and Excel workbook generation."""

from __future__ import annotations

import re
from datetime import datetime
from io import BytesIO
from pathlib import Path

import duckdb
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from src.storage import DEFAULT_DB_PATH, PROJECT_ROOT

DEFAULT_REPORT_SQL_PATH = PROJECT_ROOT / "sql" / "daily_report.sql"
REPORT_SHEETS = {
    "data_overview": "数据概览",
    "issue_type_statistics": "问题类型统计",
    "region_statistics": "区域统计",
    "high_priority_issues": "高优先级问题",
    "error_details": "错误数据明细",
}
_QUERY_MARKER = re.compile(r"^--\s*name:\s*([a-z][a-z0-9_]*)\s*$", re.MULTILINE)


class DailyReportError(RuntimeError):
    """Raised when report SQL or workbook generation is invalid."""


def load_report_queries(sql_path: Path = DEFAULT_REPORT_SQL_PATH) -> dict[str, str]:
    """Load named statements from sql/daily_report.sql."""
    path = Path(sql_path)
    if not path.is_file():
        raise DailyReportError(f"日报 SQL 文件不存在：{path}")
    content = path.read_text(encoding="utf-8")
    matches = list(_QUERY_MARKER.finditer(content))
    queries: dict[str, str] = {}
    for position, match in enumerate(matches):
        start = match.end()
        end = matches[position + 1].start() if position + 1 < len(matches) else len(content)
        statement = content[start:end].strip()
        if statement:
            queries[match.group(1)] = statement
    missing = [name for name in REPORT_SHEETS if name not in queries]
    if missing:
        raise DailyReportError(f"日报 SQL 缺少查询：{', '.join(missing)}")
    return queries


def query_daily_report(
    batch_id: str,
    database_path: Path = DEFAULT_DB_PATH,
    sql_path: Path = DEFAULT_REPORT_SQL_PATH,
) -> dict[str, pd.DataFrame]:
    """Execute every core report metric in SQL for one upload batch."""
    queries = load_report_queries(sql_path)
    connection = duckdb.connect(str(Path(database_path)), read_only=True)
    try:
        return {
            name: connection.execute(query, [batch_id] * query.count("?")).fetchdf()
            for name, query in queries.items()
            if name in REPORT_SHEETS
        }
    except Exception as exc:
        raise DailyReportError(f"日报查询失败：{exc}") from exc
    finally:
        connection.close()


def daily_report_name(generated_at: datetime) -> Path:
    return Path(f"road_quality_daily_{generated_at:%Y%m%d}.xlsx")


def _format_worksheet(worksheet) -> None:
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    worksheet.row_dimensions[1].height = 24
    for cell in worksheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for column_cells in worksheet.columns:
        width = min(max((len(str(cell.value or "")) for cell in column_cells), default=0) + 2, 55)
        worksheet.column_dimensions[get_column_letter(column_cells[0].column)].width = max(width, 10)


def build_daily_report_excel(report_tables: dict[str, pd.DataFrame]) -> bytes:
    """Create and reopen a five-sheet XLSX report from SQL result tables."""
    missing = [name for name in REPORT_SHEETS if name not in report_tables]
    if missing:
        raise DailyReportError(f"日报数据缺少工作表：{', '.join(missing)}")
    output = BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        for query_name, sheet_name in REPORT_SHEETS.items():
            report_tables[query_name].to_excel(writer, sheet_name=sheet_name, index=False)
        for worksheet in writer.book.worksheets:
            _format_worksheet(worksheet)

    workbook_bytes = output.getvalue()
    workbook = load_workbook(BytesIO(workbook_bytes), read_only=True, data_only=True)
    try:
        if workbook.sheetnames != list(REPORT_SHEETS.values()):
            raise DailyReportError("日报工作表校验失败。")
    finally:
        workbook.close()
    return workbook_bytes
